#!/usr/bin/env python3
import json
import os
import tempfile
from pathlib import Path
import openpyxl
from content_export import export

result={'mode':'rewrite','title':'既存のタイトル','audience':'想定読者','related_keywords':[],
        'reference':{'url':'https://example.com/article','fetched_at':'2026-10-20T01:00:00Z'},
        'lead_guidance':['既存の見出しを維持する。','実測AI回答に基づく改善点をコメントする。'],
        'sections':[
            dict(level='H2',heading='元の見出し',word_count=0,points=['子の見出しを紹介。'],format='',sources=['https://example.com/article#intro'],rewrite_comments=['（表の作成）条件の比較表を作成']),
            dict(level='H4',heading='細かな見出し',word_count=60,points=['既存本文を説明。'],format='',sources=['https://example.com/article#detail'],rewrite_comments=[]),
            dict(level='H3',heading='お問い合わせ',word_count=30,points=['問い合わせ方法を紹介。'],format='',sources=['https://example.com/contact'],rewrite_comments=['（可能なら）担当者への取材で事例を加える。'])],
        'competitors':[],'fanout':[],'limitations':[]}
payload={'job':{'mode':'rewrite'},'items':[dict(keyword='SEO',status='completed',result=result,evidence={},ai_prompts={'prompts':['SEOとは？']})]}
with tempfile.TemporaryDirectory() as directory:
    output=Path(directory)/'rewrite.xlsx'
    export(payload,output)
    sheet=openpyxl.load_workbook(output).active
    headers=[row for row in sheet.iter_rows() if str(row[1].value or '').startswith(('見出し(','小見出し('))]
    assert [row[2].value for row in headers]==['元の見出し','細かな見出し','お問い合わせ']
    assert headers[1][1].value=='小見出し(H4)'
    assert sheet['C4'].value==3 and sheet['B4'].value=='既存見出し数'
    assert sheet['C3'].value=='https://example.com/article' and sheet['C11'].value=='リライト（LLMO調査）'
    assert all(row[3].value is None and row[5].value=='LLMO改善コメント' for row in headers)
    assert all(row[5].border.bottom.style=='thin' and row[5].fill.fgColor.rgb==headers[0][5].fill.fgColor.rgb for row in headers)
    texts=[cell.value for row in sheet for cell in row]
    assert not any(str(text).startswith('まとめ(') for text in texts)
    assert all(cell.value is None for row in sheet.iter_rows(min_col=7,max_col=16) for cell in row)
    red=[cell for row in sheet for cell in row if str(cell.value or '').startswith(('（可能なら）','（表の作成）'))]
    assert len(red)==2 and all(cell.column==6 and cell.font.color.rgb=='FFFF0000' for cell in red)
    assert all(cell.border.bottom.style=='thin' for cell in red)
    assert all(cell.alignment.wrap_text for cell in red)
    assert all(cell.comment is None for row in sheet for cell in row)
    last=max(cell.row for row in sheet for cell in row if cell.value)
    assert all(sheet.cell(last,c).border.bottom.style=='thin' for c in range(2,7))
    audit=json.loads(output.with_suffix('.coverage.json').read_text())
    assert audit[0]['heading_count']==3 and not any('競合' in value for value in audit[0]['unavailable'])
    if os.environ.get('CONTENT_REWRITE_TEST_OUTPUT'): export(payload,Path(os.environ['CONTENT_REWRITE_TEST_OUTPUT']))
print('PASS: rewrite Excel preserves headings/levels, separate comments, zero-count parents, no new outline/competitor tables')
