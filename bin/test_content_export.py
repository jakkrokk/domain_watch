#!/usr/bin/env python3
import json
import os
import tempfile
from copy import copy, deepcopy
from pathlib import Path
import openpyxl
from content_export import export, sentences, related_keywords, ROOT

sample = openpyxl.load_workbook(ROOT / 'resources/content-generation/template.xlsx')['サンプル']
result = dict(title='土地に合う活用を選ぶ', related_keywords=['土地の適性'], audience='土地所有者',
              lead='旧形式の完成原稿。執筆指示として出力しない。', assumptions=[], limitations=[], llmo_analysis='',
              sections=[dict(level='H2',heading='土地の条件を確認する',word_count=200,
                             points=['接道を確認する。需要も調べる。','・人口\n・交通量'],format='',
                             sources=['https://example.com/a','https://example.com/b']),
                        dict(level='H3',heading='失敗要因を調べる',word_count=300,points=['管理費を確認する。'],format='',sources=['https://example.com/c'])],
              competitors=[], fanout=[])
keywords = ['田舎 土地活用 成功例', '田舎 土地活用 成功例', 'A/B:*?[]', '長いテーマ'*10, 'テンプレ']
items = [dict(position=i+1, keyword=kw, status='completed', result=deepcopy(result), evidence={}) for i,kw in enumerate(keywords)]
items[1]['result']['lead_guidance'] = ['読者の不安に共感し、成功条件を示す。','土地の適性を判断できるメリットを伝える。']
manual_prompts = ['土地活用の成功条件は？具体例も教えてください。', '=未確認の数値は使わず、失敗要因を教えてください。']
inferred_prompts = ['田舎の土地の適性をどう判断しますか？', '土地活用の需要はどう調べますか？', '活用方法を選ぶ際の注意点は何ですか？']
items[1]['ai_prompts'] = json.dumps(dict(origin='manual', prompts=manual_prompts), ensure_ascii=False)
items[2]['ai_prompts'] = dict(origin='inferred', prompts=inferred_prompts)
with tempfile.TemporaryDirectory() as directory:
    output = Path(directory)/'test.xlsx'
    export(dict(job={'notes':'判断基準を重視する。'},items=items),output)
    wb = openpyxl.load_workbook(output)
    assert len(wb.worksheets) == 5
    values = openpyxl.load_workbook(output, data_only=True)
    assert values.worksheets[0]['C4'].value == 700
    assert wb.worksheets[0]['C4'].value == '=SUM(D:D)'
    report = json.loads(output.with_suffix('.coverage.json').read_text())
    assert len(report) == 5 and report[0]['competitor_count'] == 0
    assert report[0]['unavailable']

    assert wb.sheetnames[0] == keywords[0] and wb.sheetnames[1].endswith('(2)')
    assert len(set(n.casefold() for n in wb.sheetnames)) == 5
    assert all(len(n)<=31 and not any(c in n for c in '\\/*?:[]') for n in wb.sheetnames)
    for s in wb:
        assert s.freeze_panes is None
        assert s.sheet_view.showGridLines == sample.sheet_view.showGridLines
        assert s.sheet_view.zoomScale == sample.sheet_view.zoomScale
        for label in ['H9','N9','O10','P10','H23','I24','J24','K24']:
            assert s[label].value == sample[label].value, label
        assert all(s[f'I{r}'].value for r in range(11,21))
        assert all(s[f'I{r}'].value for r in range(25,75))
        assert all(c.comment is None for row in s for c in row)
        assert s['E16'].alignment.horizontal == 'fill'
        assert not any(c.has_style for row in s.iter_rows(min_row=23,max_col=6) for c in row)

        assert all(s.cell(10,c).value is None for c in range(2,7))
        assert s.row_dimensions[10].height == sample.row_dimensions[7].height
        assert s['C11'].value == 'リード文' and s['B12'].value == '概要'
        assert s['C12'].value and s['C13'].value and s['C14'].value is None
        assert '旧形式' not in s['C12'].value
        assert [s[f'C{r}'].value for r in range(16,20)] == ['接道を確認する。','需要も調べる。','・人口','・交通量']
        assert s['C20'].value is None and s['B21'].value == '小見出し(H3)'
        assert s['E15'].value == '参照URL' and s['E16'].value == 'https://example.com/a' and s['E17'].value == 'https://example.com/b'
        assert s['E21'].value == '参照URL' and s['E22'].value == 'https://example.com/c'
        assert s['F15'].value == 'クライアント様コメント' and s['F21'].value is None
        for cell, ref in [('C11','C10'),('C12','C11'),('E15','E14'),('E16','E15'),('E21','E19'),('C21','C19')]:
            for attr in ['font','fill','border','alignment','number_format']:
                if cell == 'E16' and attr == 'alignment': continue
                assert copy(getattr(s[cell],attr)) == copy(getattr(sample[ref],attr)), (cell,attr)
        assert s.row_dimensions[11].height == sample.row_dimensions[10].height
        assert s.column_dimensions['E'].width == sample.column_dimensions['E'].width
        assert s['C2'].comment is None
        assert all('\n' not in str(s[f'{c}{r}'].value or '') for c in ['C','E'] for r in range(10,24))
    assert wb.worksheets[1]['C12'].value == items[1]['result']['lead_guidance'][0]
    assert wb.worksheets[0]['B9'].value == '想定プロンプト' and wb.worksheets[0]['C9'].value == '未取得'
    assert wb.worksheets[1]['C9'].value == '\n'.join('・' + prompt for prompt in manual_prompts)
    assert wb.worksheets[2]['C9'].value == '\n'.join('・' + prompt for prompt in inferred_prompts)
    assert wb.worksheets[1]['C9'].data_type == 's' and wb.worksheets[1]['C9'].alignment.wrap_text
    assert wb.worksheets[1]['B9']._style == sample['B8']._style
    assert wb.worksheets[1]['C9']._style == sample['C8']._style
    assert wb.worksheets[2].row_dimensions[9].height >= sample.row_dimensions[8].height
    filled = deepcopy(items[0])
    filled['result']['related_keywords'] = [f'土地活用 条件{i}' for i in range(50)]
    filled['result']['competitors'] = [dict(rank=0,title=f'確認した競合{i}',url=f'https://example.com/{i}',word_count=0,headings=['土地の条件','注意点']) for i in range(10)]
    filled['result']['lead_guidance'] = ['田舎の土地活用に悩む読者に、土地の条件から判断する方法を示す。','成功条件と失敗要因を比較して自分の土地に合う方法を選べることを伝える。']
    full_output = Path(directory)/'full.xlsx'
    export(dict(job={},items=[filled]),full_output)
    full = openpyxl.load_workbook(full_output).worksheets[0]
    assert [full[f'H{r}'].value for r in range(11,21)] == list(range(1,11))
    assert [full[f'N{r}'].value for r in range(11,31)] == [n for n in range(1,11) for _ in range(2)]
    assert full['I20'].value == '確認した競合9'
    assert full['J20'].value == 'https://example.com/9'
    assert full['I74'].value == '土地活用 条件49'
    assert full['O30'].value == '注意点' and full['P30'].value == 'https://example.com/9'
    assert all(full[f'I{r}'].value != '未取得' for r in list(range(11,21))+list(range(25,75)))
    filled['result']['sections'][0]['format'] = '可食部を扱う際の箇条書き'
    filled['result']['sections'][0]['llmo_reasons'] = ['AI回答の比較表を参考に、土地条件を比較できるようにする。']
    export(dict(job={},items=[filled]),full_output)
    revised = openpyxl.load_workbook(full_output).active
    texts = [revised[f'C{r}'].value for r in range(14,25)]
    instruction = '上記の要点や確認事項を箇条書きで示す。'
    reason = '※LLMO用：AI回答の比較表を参考に、土地条件を比較できるようにする。'
    assert instruction in texts and reason in texts and texts.index(reason) > texts.index(instruction)
    assert all(not str(c.value).startswith(('表現：','形式：')) and c.comment is None for row in revised for c in row)
    complete_report = json.loads(full_output.with_suffix('.coverage.json').read_text())[0]
    assert complete_report['competitor_count'] == 10 and complete_report['related_keyword_count'] == 50

    attributed = deepcopy(filled)
    attributed['evidence'] = {
        'openai': {'citation_urls': ['https://example.com/a'], 'source_urls': ['https://example.com/reference-only']},
        'gemini': {'citation_urls': [], 'source_urls': []},
    }
    attributed['result']['sections'][0]['source_attributions'] = [
        dict(source_name='自治体の土地活用ガイド', source_url='https://example.com/a',
             supported_claim='用途地域によって建築できる建物が異なる。',
             writing_guidance='自治体の土地活用ガイドによると用途地域の確認が必要であると、出典名を本文に示して説明する。',
             origins=['openai_citation']),
        dict(source_name='地域の需要調査', source_url='https://example.com/original-study',
             supported_claim='調査対象地域の需要を示す資料であり、全国に一般化しない。',
             writing_guidance='地域の需要調査を出典として、調査時期と対象地域を添えて説明する。',
             origins=['additional_research']),
    ]
    attributed['result']['sections'][1]['source_attributions'] = []
    attributed_output = Path(directory)/'attributed.xlsx'
    export(dict(job={}, items=[attributed]), attributed_output)
    attributed_sheet = openpyxl.load_workbook(attributed_output).active
    outline = [str(row[0].value or '') for row in attributed_sheet.iter_rows(min_col=3, max_col=3)]
    for attribution in attributed['result']['sections'][0]['source_attributions']:
        assert 'LLMOのため出典を明示（' + attribution['source_name'] + '）' in outline
    assert not any(text.startswith(('根拠として','出典を示す執筆指示','LLMOとして入れた理由')) for text in outline)
    urls = [row[0].value for row in attributed_sheet.iter_rows(min_col=5, max_col=5)]
    assert urls.count('https://example.com/a') == 1
    assert 'https://example.com/original-study' in urls
    assert 'https://example.com/reference-only' not in urls  # Never export all AI candidates as accepted evidence.
    assert outline.index('LLMOのため出典を明示（地域の需要調査）') < outline.index(reason)
    assert outline.count('失敗要因を調べる') == 1
    assert all(cell.comment is None for row in attributed_sheet for cell in row)
    assert all('\n' not in str(row[0].value or '') for row in attributed_sheet.iter_rows(min_row=10, min_col=3, max_col=3))
    attributed['ai_prompts'] = dict(origin='inferred', prompts=inferred_prompts)
    attributed['result']['keyword_suggestions'] = dict(status='completed', values=['土地活用 成功例','土地活用 費用','土地活用 失敗'])
    attributed['result']['related_keywords'] = ['site:example.com data', 'https://example.com', '土地活用 費用','土地活用 成功例','無関係な候補']
    attributed['evidence']['openai']['search_queries'] = ['site:example.com research']
    attributed['result']['sections'][0]['format'] = '（表の作成）土地条件の比較表を作成'
    attributed['result']['sections'][1]['points'] += ['（可能なら）営業担当への取材で、実際の相談例を加える。']
    attributed['result']['sections'].append(dict(level='H2',heading='まとめ：土地に合う方法を選ぶ',word_count=100,
        points=['土地の条件と需要を確認する。','取り組みやすい方法を選ぶ。'],format='',sources=['https://example.com/summary']))
    export(dict(job={},items=[attributed]),attributed_output)
    checked = openpyxl.load_workbook(attributed_output).active
    assert checked['C3'].value == '土地活用 費用、土地活用 成功例'
    assert checked['I25'].value == '土地活用 成功例'
    assert 'site:example.com research' not in [checked[f'I{r}'].value for r in range(25,75)]
    optional = next(c for row in checked.iter_rows(min_col=3,max_col=3) for c in row if str(c.value).startswith('（可能なら）'))
    table = next(c for row in checked.iter_rows(min_col=3,max_col=3) for c in row if str(c.value).startswith('（表の作成）'))
    assert optional.font.color.rgb == 'FFFF0000' and table.font.color.rgb == 'FFFF0000'
    last = max(r for r in range(15,checked.max_row+1) if checked.cell(r,3).value)
    assert all(checked.cell(last,col).border.bottom.style == 'thin' for col in range(2,7))
    assert checked.cell(last,6).border.right.style == 'thin'
    if os.environ.get('CONTENT_EXPORT_TEST_OUTPUT'):
        export(dict(job={},items=[attributed]),Path(os.environ['CONTENT_EXPORT_TEST_OUTPUT']))
assert sentences('「失敗した？」と聞く。条件を調べる。') == ['「失敗した？」と聞く。','条件を調べる。']
assert related_keywords(dict(related_keywords=['site:example.com SEO','SEO 費用','SEO 費用','https://example.com'])) == ['SEO 費用']
print('PASS: legacy and attributed outlines, source instructions and URLs, 5 themes/5 sheets, sample styles')
